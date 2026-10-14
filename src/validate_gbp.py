"""Sanity check: compare the GBP SONIA swap data with the Bank of England's published OIS spot curve.
Differences are expected to be small but non-zero: the BoE curve is a fitted zero-coupon curve and the
swap data are par rates, possibly marked at a different time of day. Writes output/validation.json."""
import io
import json
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd
import requests

ROOT = Path(__file__).resolve().parents[1]
URL = "https://www.bankofengland.co.uk/-/media/boe/files/statistics/yield-curves/oisddata.zip"
TEN = [2, 5, 10, 25]


def boe_curve():
    z = zipfile.ZipFile(io.BytesIO(requests.get(URL, timeout=120, headers={"User-Agent": "Mozilla/5.0"}).content))
    frames = []
    for name in z.namelist():
        wb = openpyxl.load_workbook(io.BytesIO(z.read(name)), read_only=True, data_only=True)
        if "4. spot curve" not in wb.sheetnames:
            continue
        rows = list(wb["4. spot curve"].iter_rows(values_only=True))
        h = next(i for i, r in enumerate(rows) if isinstance(r[0], str) and r[0].strip().startswith("years"))
        mats = list(rows[h][1:])
        recs = [{"date": pd.Timestamp(r[0]), **{f"{t}y": r[1 + mats.index(t)] for t in TEN}}
                for r in rows[h + 1:] if hasattr(r[0], "year")]
        frames.append(pd.DataFrame(recs).set_index("date"))
    return pd.concat(frames).sort_index().apply(pd.to_numeric, errors="coerce").dropna()


def main():
    g = pd.read_csv(ROOT / "data" / "gbp_swap_curve.csv", index_col=0, parse_dates=True)
    b = boe_curve()
    j = g[[f"{t}y" for t in TEN]].join(b, rsuffix="_boe", how="inner")
    out = {"overlap_start": str(j.index[0].date()), "overlap_end": str(j.index[-1].date()), "days": len(j)}
    for t in TEN:
        a, c = j[f"{t}y"] * 100, j[f"{t}y_boe"] * 100
        out[f"{t}y"] = {"mean_level_diff_bp": round(float((a - c).mean()), 1), "sd_level_diff_bp": round(float((a - c).std()), 1),
                        "daily_change_corr": round(float(a.diff().corr(c.diff())), 3)}
    (ROOT / "output" / "validation.json").write_text(json.dumps(out, indent=2))
    print(json.dumps(out, indent=2))


if __name__ == "__main__":
    main()
